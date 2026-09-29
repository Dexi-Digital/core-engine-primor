"""Leitura estrita do Relatório Completo, sem preencher rateios por inferência."""

from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ABA = "Relatório Completo"
MAX_BYTES = 30 * 1024 * 1024
MAX_LINHAS = 200_000
CAMPOS = {
    "EMPRESA": "empresa",
    "CLIENTE / FORNECEDOR": "contraparte",
    "TÍTULO": "titulo",
    "DATA PAGAMENTO": "data_pagamento",
    "PERÍODO": "periodo",
    "DATA INCLUSÃO NF": "data_inclusao_nf",
    "DATA EMISSÃO": "data_emissao",
    "VALOR ORIGINAL": "valor_original",
    "VALOR INCREMENTO": "valor_incremento",
    "VALOR ABATIMENTO": "valor_abatimento",
    "VALOR AMORTIZADO": "valor_amortizado",
    "VALOR ACRÉSCIMO": "valor_acrescimo",
    "VALOR DEDUÇÃO": "valor_deducao",
    "VALOR PAGO": "valor_pago",
    "VALOR ATUAL": "valor_atual",
    "LOCAL": "local",
    "CENTRO DE CUSTO": "centro_custo",
    "ID NATUREZA": "id_natureza",
    "I": "natureza_i",
    "II": "natureza_ii",
    "III": "natureza_iii",
    "NATUREZA": "natureza",
    "VALOR APROPRIADO": "valor_apropriado",
}
DATAS = {"data_pagamento", "periodo", "data_inclusao_nf", "data_emissao"}


def classificar(i: str | None, ii: str | None, iii: str | None) -> str | None:
    """Regras específicas de aportes/equipamentos precedem as gerais."""

    def nivel(value: str | None) -> int | None:
        return int(value) if value and value.isdigit() else None

    a, b, c = nivel(i), nivel(ii), nivel(iii)
    if a == 3:
        return "Investimento"
    if a == 1:
        if b == 1 and c is None:
            return None
        return {(1, 3): "Devolução de Aporte de SCP entrada", (1, 8): "Aporte entrada"}.get(
            (b, c), "Receita"
        )
    if a != 2:
        return None
    if b == 3:
        if c is None:
            return None
        if c in (16, 17, 21):
            return "Locação de Equipamento"
        if c == 27:
            return "Equipamento Próprio"
        return "Material"
    if b == 9:
        return {1: "Aporte saída", 2: "Devolução de Aporte de SCP saída"}.get(c)
    return {
        1: "Execução e Escritório Local",
        2: "Mão de Obra",
        4: "Material",
        5: "Imposto",
        6: "Imposto",
        7: "Imposto",
        8: "Serviço",
    }.get(b)


def _data(value: object, epoch: datetime) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        parsed = from_excel(value, epoch)
        if isinstance(parsed, datetime):
            return parsed.date()
    if isinstance(value, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError("data inválida")


def _valor(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    text = str(value).strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        number = Decimal(text)
        if not number.is_finite() or abs(number) >= Decimal("1e16"):
            raise ValueError("valor fora do limite")
        return number.quantize(Decimal("0.00000001"))
    except InvalidOperation as exc:
        raise ValueError("valor numérico inválido") from exc


def _texto(cell) -> str | None:
    value = cell.value
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if int(value) != value:
            raise ValueError("código numérico fracionário")
        value = str(int(value))
        if cell.number_format and set(cell.number_format) == {"0"}:
            value = value.zfill(len(cell.number_format))
    return str(value).strip() or None


def _registro(cells: dict, epoch: datetime) -> tuple[dict, list[str]]:
    result = {}
    raw = {}
    for header, campo in CAMPOS.items():
        cell = cells[header]
        value = cell.value
        raw[header] = value.isoformat() if isinstance(value, (date, datetime)) else value
        if cell.data_type in ("f", "e"):
            raise ValueError(f"{header}: fórmula ou erro de Excel; exporte os valores")
        try:
            if campo in DATAS:
                parsed = _data(value, epoch)
            elif campo.startswith("valor_"):
                parsed = _valor(value)
            else:
                parsed = _texto(cell)
                limit = 32 if campo.startswith("natureza_") else 64 if campo == "empresa" else 255
                if campo not in ("contraparte", "natureza") and parsed and len(parsed) > limit:
                    raise ValueError(f"texto excede {limit} caracteres")
            result[campo] = parsed
        except (ValueError, OverflowError) as exc:
            raise ValueError(f"{header}: {exc}") from exc
    for required in ("empresa", "titulo", "valor_apropriado"):
        if result[required] is None:
            raise ValueError(f"{required}: preenchimento obrigatório")
    result["raw"] = raw
    emissao = result["data_emissao"]
    result["competencia"] = emissao.replace(day=1) if emissao else None
    result["classificacao"] = classificar(
        result["natureza_i"], result["natureza_ii"], result["natureza_iii"]
    )
    avisos = []
    if not emissao:
        avisos.append("Sem data de emissão; competência pendente")
    if not result["classificacao"]:
        avisos.append("Natureza sem classificação; dados originais preservados")
    return result, avisos


def ler_planilha(content: bytes) -> Iterator[tuple[int, dict | None, list[str]]]:
    """Produz registros/erros por linha; erro de estrutura rejeita o arquivo inteiro."""
    if len(content) > MAX_BYTES:
        raise ValueError("Arquivo excede 30 MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 512 * 1024 * 1024:
                raise ValueError("Conteúdo descompactado excede 512 MB")
    except BadZipFile as exc:
        raise ValueError("Arquivo não é um XLSX válido") from exc
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if ABA not in workbook.sheetnames:
            raise ValueError(f"Aba obrigatória: {ABA}")
        sheet = workbook[ABA]
        if (sheet.max_row or 0) > MAX_LINHAS + 1 or (sheet.max_column or 0) > 100:
            raise ValueError("Planilha excede os limites de 200.000 linhas / 100 colunas")
        rows = sheet.iter_rows()
        header = next(rows, ())
        names = [str(cell.value).strip() if cell.value is not None else "" for cell in header]
        while names and not names[-1]:
            names.pop()
        if len(names) != len(CAMPOS) or set(names) != set(CAMPOS):
            raise ValueError("Cabeçalho deve conter exatamente as 23 colunas do Relatório Completo")
        for number, row in enumerate(rows, 2):
            if number > MAX_LINHAS + 1:
                raise ValueError("Planilha excede 200.000 linhas")
            if all(cell.value is None for cell in row):
                continue
            try:
                if any(cell.value is not None for cell in row[len(names) :]):
                    raise ValueError("Valores em colunas sem cabeçalho")
                record, warnings = _registro(dict(zip(names, row, strict=False)), workbook.epoch)
                yield number, record, warnings
            except ValueError as exc:
                yield number, None, [str(exc)]
    finally:
        workbook.close()
