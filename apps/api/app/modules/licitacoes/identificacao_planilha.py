"""Identificacao da planilha orcamentaria editavel (Captador, secao 8).

Combina tres criterios para reduzir falso positivo -- a identificacao
NAO depende so do nome do arquivo, porque muitos portais publicam
nomenclatura generica:

  1. extensao: apenas XLSX / XLS / ODS sao elegiveis (PDF e arquivado,
     mas nao substitui a planilha editavel no aceite do cliente);
  2. nome do arquivo: termos do vocabulario de orcamento (10 pts cada);
  3. conteudo interno: cabecalhos tipicos de planilha de orcamento
     nas primeiras linhas (5 pts cada termo encontrado).

Funcoes puras, sem IO/DB -- quem le bytes do storage e o caller
(`processamento.py`). Bytes corrompidos pontuam 0 no conteudo em vez
de propagar excecao: um zip invalido nao pode derrubar o pipeline.
"""
from __future__ import annotations

import io
import logging
import unicodedata
from pathlib import PurePosixPath

logger = logging.getLogger(__name__)

EXTENSOES_ELEGIVEIS: frozenset[str] = frozenset({".xlsx", ".xls", ".ods"})

# Planilha principal exige score >= LIMIAR_PRINCIPAL: um termo no nome
# OU dois no conteudo. Evita eleger um xlsx qualquer (lista de presenca)
# como orcamento.
LIMIAR_PRINCIPAL: int = 10

# Vocabulario inicial do Projeto Tecnico (secao 8.1). Radicais sem
# acento -- a comparacao roda sobre texto normalizado.
TERMOS_NOME: tuple[str, ...] = (
    "orcament",
    "planilha",
    "composic",
    "cronograma",
    "bdi",
    "quantitativo",
    "preco unitario",
    "memoria de calculo",
    "custo",
    "valor total",
)

TERMOS_CONTEUDO: tuple[str, ...] = (
    "item",
    "unid",
    "quant",
    "preco unitario",
    "valor unitario",
    "bdi",
    "total",
    "descricao",
)

_PONTOS_NOME = 10
_PONTOS_CONTEUDO = 5

# Limites de varredura do conteudo: suficiente para achar o cabecalho,
# barato o bastante para rodar em serverless.
_MAX_ABAS = 3
_MAX_LINHAS = 30
_MAX_COLUNAS = 20


def normalizar(texto: str) -> str:
    """Minusculas, sem acento; `_`, `-` e quebras viram espaco unico."""
    decomposto = unicodedata.normalize("NFKD", texto.lower())
    sem_acento = "".join(c for c in decomposto if not unicodedata.combining(c))
    for sep in ("_", "-", "\n", "\t"):
        sem_acento = sem_acento.replace(sep, " ")
    return " ".join(sem_acento.split())


def score_nome(filename: str) -> int:
    nome = normalizar(filename)
    return sum(_PONTOS_NOME for termo in TERMOS_NOME if termo in nome)


def _texto_celulas_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pedacos: list[str] = []
    try:
        for ws in wb.worksheets[:_MAX_ABAS]:
            for row in ws.iter_rows(
                min_row=1, max_row=_MAX_LINHAS, max_col=_MAX_COLUNAS, values_only=True
            ):
                pedacos.extend(str(v) for v in row if v is not None)
    finally:
        wb.close()
    return " ".join(pedacos)


def _texto_celulas_ods(data: bytes) -> str:
    from odf import teletype
    from odf.opendocument import load as ods_load
    from odf.table import Table, TableCell, TableRow

    doc = ods_load(io.BytesIO(data))
    pedacos: list[str] = []
    for table in doc.spreadsheet.getElementsByType(Table)[:_MAX_ABAS]:
        for row in table.getElementsByType(TableRow)[:_MAX_LINHAS]:
            for cell in row.getElementsByType(TableCell)[:_MAX_COLUNAS]:
                texto = teletype.extractText(cell)
                if texto:
                    pedacos.append(texto)
    return " ".join(pedacos)


def score_conteudo(filename: str, data: bytes) -> int:
    """Pontua cabecalhos tipicos nas primeiras celulas. 0 se ilegivel."""
    ext = PurePosixPath(filename).suffix.lower()
    try:
        if ext == ".xlsx":
            texto = _texto_celulas_xlsx(data)
        elif ext == ".ods":
            texto = _texto_celulas_ods(data)
        else:
            # .xls legado: openpyxl nao le. Score de conteudo 0 --
            # decisao registrada no plano (nome + extensao bastam).
            return 0
    except Exception as exc:  # noqa: BLE001 -- qualquer corrupcao = score 0
        logger.warning("planilha ilegivel (%s): %s", filename, exc)
        return 0
    texto_norm = normalizar(texto)
    return sum(
        _PONTOS_CONTEUDO for termo in TERMOS_CONTEUDO if termo in texto_norm
    )


def classificar_anexo(filename: str, data: bytes | None) -> int | None:
    """Score total do anexo, ou `None` se a extensao e inelegivel."""
    ext = PurePosixPath(filename).suffix.lower()
    if ext not in EXTENSOES_ELEGIVEIS:
        return None
    total = score_nome(filename)
    if data is not None:
        total += score_conteudo(filename, data)
    return total
