"""Importação por versão: rateios, auditoria, autorização e exportação reais em SQLite."""

import hashlib
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.audit.models import AuditLog
from app.main import app
from app.modules.financeiro_importacao import router, service
from app.modules.financeiro_importacao.models import ApropriacaoFinanceira, ImportacaoFinanceira
from app.modules.financeiro_importacao.parser import ABA, CAMPOS, classificar, ler_planilha
from app.modules.financeiro_importacao.schemas import Filtros
from app.modules.licitacoes.storage import LocalStorage


def planilha(*changes, aba=ABA):
    wb = Workbook()
    ws = wb.active
    ws.title = aba
    ws.append(list(CAMPOS))
    defaults = {
        "EMPRESA": "EMP",
        "TÍTULO": "00012",
        "CLIENTE / FORNECEDOR": "Fornecedor",
        "DATA EMISSÃO": date(2026, 6, 22),
        "VALOR ORIGINAL": 100,
        "VALOR PAGO": 100,
        "LOCAL": "B001",
        "CENTRO DE CUSTO": "01.01",
        "ID NATUREZA": "2.03.018",
        "I": "2",
        "II": "03",
        "III": "018",
        "NATUREZA": "COMBUSTÍVEIS",
        "VALOR APROPRIADO": 100,
    }
    for change in changes or ({},):
        row = defaults | change
        ws.append([row.get(header) for header in CAMPOS])
    content = BytesIO()
    wb.save(content)
    return content.getvalue()


@pytest.mark.parametrize(
    ("niveis", "expected"),
    [
        (("1", "01", "003"), "Devolução de Aporte de SCP entrada"),
        (("1", "01", "008"), "Aporte entrada"),
        (("1", "02", "001"), "Receita"),
        (("1", None, None), "Receita"),
        (("2", "02", None), "Mão de Obra"),
        *[(("2", "03", n), "Locação de Equipamento") for n in ("016", "017", "021")],
        (("2", "03", "027"), "Equipamento Próprio"),
        (("2", "03", "018"), "Material"),
        (("2", "04", "027"), "Material"),
        (("2", "01", None), "Execução e Escritório Local"),
        *[(("2", n, None), "Imposto") for n in ("05", "06", "07")],
        (("2", "08", None), "Serviço"),
        (("2", "09", "001"), "Aporte saída"),
        (("2", "09", "002"), "Devolução de Aporte de SCP saída"),
        (("3", None, None), "Investimento"),
        (("2", "03", None), None),
        (("1", "01", None), None),
        (("2", "09", "005"), None),
        ((None, None, None), None),
    ],
)
def test_classificacoes(niveis, expected):
    assert classificar(*niveis) == expected


def test_preserva_rateios_zeros_datas_e_valores_vazios():
    rows = list(
        ler_planilha(
            planilha(
                {"VALOR APROPRIADO": "30,25"},
                {"VALOR ORIGINAL": None, "VALOR PAGO": None, "VALOR APROPRIADO": "69,75"},
            )
        )
    )
    assert len(rows) == 2
    assert rows[0][1]["competencia"] == date(2026, 6, 1)
    assert rows[0][1]["titulo"] == "00012"
    assert rows[0][1]["natureza_ii"] == "03"
    assert rows[1][1]["valor_original"] is None
    assert sum(r[1]["valor_apropriado"] for r in rows) == Decimal("100")


@pytest.mark.parametrize(
    "change",
    [
        {"VALOR APROPRIADO": "NaN"},
        {"VALOR APROPRIADO": "Infinity"},
        {"VALOR APROPRIADO": None},
        {"DATA EMISSÃO": "31/02/2026"},
        {"VALOR ORIGINAL": "=SUM(1,2)"},
        {"EMPRESA": None},
    ],
)
def test_linha_invalida_e_visivel(change):
    row = list(ler_planilha(planilha(change)))[0]
    assert row[0] == 2 and row[1] is None and row[2]


def test_recusa_resumo():
    with pytest.raises(ValueError, match="Aba obrigatória"):
        list(ler_planilha(planilha(aba="Relatório Final (Resumo)")))


@pytest.fixture
def storage(tmp_path):
    instance = LocalStorage(tmp_path)

    async def override():
        yield instance

    app.dependency_overrides[router.get_storage] = override
    yield instance
    app.dependency_overrides.pop(router.get_storage, None)


@pytest.fixture
def dispatcher(monkeypatch):
    class Fake:
        calls = []

        def send_task(self, name, **kwargs):
            self.calls.append((name, kwargs))

    fake = Fake()
    monkeypatch.setattr(service, "get_celery_dispatcher", lambda: fake)
    return fake


async def preparar_e_enviar(client, headers, content):
    prepared = await client.post(
        "/api/v1/financeiro/importacoes",
        headers=headers,
        json={"nome_arquivo": "base.xlsx", "sha256": hashlib.sha256(content).hexdigest()},
    )
    assert prepared.status_code == 201, prepared.text
    data = prepared.json()
    id_ = data["lote"]["id"]
    sent = await client.put(
        f"/api/v1/financeiro/importacoes/{id_}/arquivo",
        content=content,
        headers={"Authorization": f"Bearer {data['upload_token']}"},
    )
    assert sent.status_code == 202, sent.text
    return id_, data["upload_token"]


@pytest.mark.asyncio
async def test_fluxo_completo_e_idempotencia(
    api_client, auth_headers, db_session, storage, dispatcher
):
    content = planilha(
        {"VALOR APROPRIADO": 40},
        {"VALOR ORIGINAL": None, "VALOR PAGO": None, "VALOR APROPRIADO": 60},
    )
    id_, token = await preparar_e_enviar(api_client, auth_headers, content)
    assert dispatcher.calls[0][1]["queue"] == "financeiro"
    # Capacidade de upload não autentica endpoints normais nem download.
    restricted = {"Authorization": f"Bearer {token}"}
    assert (
        await api_client.get("/api/v1/financeiro/importacoes", headers=restricted)
    ).status_code == 401
    assert (
        await api_client.get(f"/api/v1/financeiro/importacoes/{id_}/arquivo", headers=restricted)
    ).status_code == 401
    await service.processar(db_session, storage, id_)
    await service.processar(db_session, storage, id_)
    assert (await db_session.scalar(select(func.count()).select_from(ApropriacaoFinanceira))) == 2
    lote = await db_session.get(ImportacaoFinanceira, id_)
    assert lote.status == "validado" and lote.erros == 0
    before = await api_client.get(
        "/api/v1/financeiro/importacoes/apropriacoes", headers=auth_headers
    )
    assert before.json()["total"] == 0
    confirm = await api_client.post(
        f"/api/v1/financeiro/importacoes/{id_}/confirmar", headers=auth_headers, json={}
    )
    assert confirm.status_code == 200, confirm.text
    active = (
        await api_client.get(
            "/api/v1/financeiro/importacoes/apropriacoes?empresa=EMP", headers=auth_headers
        )
    ).json()
    assert active["total"] == 2 and Decimal(active["valor_apropriado"]) == 100
    duplicate = await service.preparar(
        db_session, nome="renomeado.xlsx", sha256=hashlib.sha256(content).hexdigest(), actor="test"
    )
    assert duplicate.id == id_
    export_auth = await api_client.post(
        f"/api/v1/financeiro/importacoes/{id_}/download-token", headers=auth_headers
    )
    export = await api_client.get(
        f"/api/v1/financeiro/importacoes/{id_}/arquivo",
        headers={"Authorization": f"Bearer {export_auth.json()['token']}"},
    )
    assert export.status_code == 200
    wb = load_workbook(BytesIO(export.content))
    ws = wb.active
    assert ws.max_column == 25 and ws.max_row == 3
    assert ws["C2"].value == "00012" and ws["C2"].data_type == "s"
    assert ws["H3"].value is None and ws["W2"].data_type == "n"
    assert ws["G2"].is_date and ws["X2"].is_date
    logs = (
        await db_session.scalars(
            select(AuditLog).where(AuditLog.resource == "financeiro.importacao")
        )
    ).all()
    assert {log.action for log in logs} >= {"create", "validate", "confirm", "export"}


@pytest.mark.asyncio
async def test_substituicao_explicita_preserva_historico(
    api_client, auth_headers, db_session, storage, dispatcher
):
    ids = []
    for valor in (100, 200):
        id_, _ = await preparar_e_enviar(
            api_client, auth_headers, planilha({"VALOR APROPRIADO": valor})
        )
        await service.processar(db_session, storage, id_)
        ids.append(id_)
    await service.confirmar(db_session, ids[0], None, "test")
    with pytest.raises(ValueError, match="base ativa mudou"):
        await service.confirmar(db_session, ids[1], None, "test")
    await service.confirmar(db_session, ids[1], ids[0], "test")
    assert (await db_session.get(ImportacaoFinanceira, ids[0])).status == "historico"
    assert (await service.listar(db_session, ids[0], Filtros()))["total"] == 1
    assert Decimal((await service.listar(db_session, ids[1], Filtros()))["valor_apropriado"]) == 200


@pytest.mark.asyncio
async def test_erros_bloqueiam_lote_inteiro(
    api_client, auth_headers, db_session, storage, dispatcher
):
    id_, _ = await preparar_e_enviar(
        api_client, auth_headers, planilha({}, {"VALOR APROPRIADO": "inválido"})
    )
    await service.processar(db_session, storage, id_)
    lote = await db_session.get(ImportacaoFinanceira, id_)
    assert lote.erros == 1 and lote.ocorrencias[0]["linha"] == 3
    assert lote.export_path is None
    with pytest.raises(ValueError, match="sem erros"):
        await service.confirmar(db_session, id_, None, "test")


@pytest.mark.asyncio
async def test_aviso_preserva_dado_sem_inventar_classificacao(
    api_client, auth_headers, db_session, storage, dispatcher
):
    id_, _ = await preparar_e_enviar(api_client, auth_headers, planilha({"I": "9"}))
    await service.processar(db_session, storage, id_)
    lote = await db_session.get(ImportacaoFinanceira, id_)
    assert lote.avisos == 1 and lote.erros == 0 and lote.status == "validado"


@pytest.mark.asyncio
async def test_upload_exige_operador_e_hash(
    api_client, auth_headers, admin_user, db_session, storage, dispatcher
):
    content = planilha()
    req = {"nome_arquivo": "base.xlsx", "sha256": hashlib.sha256(content).hexdigest()}
    assert (await api_client.post("/api/v1/financeiro/importacoes", json=req)).status_code == 401
    admin_user.role = "leitor"
    await db_session.commit()
    assert (
        await api_client.post("/api/v1/financeiro/importacoes", json=req, headers=auth_headers)
    ).status_code == 403
    admin_user.role = "admin"
    await db_session.commit()
    prepared = (
        await api_client.post("/api/v1/financeiro/importacoes", json=req, headers=auth_headers)
    ).json()
    upload = await api_client.put(
        f"/api/v1/financeiro/importacoes/{prepared['lote']['id']}/arquivo",
        headers={"Authorization": f"Bearer {prepared['upload_token']}"},
        content=b"outro",
    )
    assert upload.status_code == 422


@pytest.mark.asyncio
async def test_falha_de_fila_e_retomada(
    api_client, auth_headers, db_session, storage, monkeypatch, dispatcher
):
    content = planilha()
    prepared = await service.preparar(
        db_session, nome="base.xlsx", sha256=hashlib.sha256(content).hexdigest(), actor="test"
    )
    prepared.storage_path = await service.salvar_arquivo(storage, prepared.id, "original", content)
    original = dispatcher.send_task

    def broken(*args, **kwargs):
        raise ConnectionError("offline")

    dispatcher.send_task = broken
    with pytest.raises(RuntimeError):
        await service.enfileirar(db_session, prepared, "test")
    assert prepared.status == "erro_fila"
    dispatcher.send_task = original
    result = await api_client.post(
        f"/api/v1/financeiro/importacoes/{prepared.id}/reprocessar", headers=auth_headers
    )
    assert result.status_code == 202 and prepared.status == "na_fila"
